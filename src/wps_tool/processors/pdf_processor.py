"""PDF 处理器。

Layer (a) TODO：下面 6 个函数体（merge/split/rotate/extract_pages/to_images/extract_text）。
PdfProcessor 类已写好当分派胶水，学生不要改它。

实现顺序提示：先 extract_pdf_text（最简，单测直接驱动），再 merge/split，最后 to_images。
每个函数的测试在 tests/test_pdf_processor.py，是 spec：实现到测试转绿即完成。

库 API 要点（pypdf 6.x / PyMuPDF 1.28，已按 venv 已装版本核对）：
- pypdf：PdfReader(path) 读、PdfWriter() 写；合并可用 writer.append(reader)（保书签）
  或循环 writer.add_page(page)；写出用 with open(out, "wb") as f: writer.write(f)；
  旋转 page.rotate(angle) 是原地修改，须把页加到新 writer；
  加密是 PdfWriter.encrypt(user_password, owner_password=None, use_128bit=True)
  ——没有 add_encryption（那是 PyPDF2 旧名）。
- fitz（PyMuPDF）：page.get_pixmap(matrix=fitz.Matrix(zoom, zoom)) 渲染；
  pix.save(path) 按扩展名定格式；page.get_text() 返回 str。
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

import fitz
from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter

from wps_tool.core.errors import UnsupportedActionError
from wps_tool.processors.base import FileProcessor
from wps_tool.utils.logging import logger

# ===== Layer (a) TODO：学生实现以下函数体 =====


def merge_pdfs(input_paths: list[str], output_path: str) -> str:
    """[学习·Layer (a) TODO] 合并多个 PDF 为一个，写出后返回 output_path。

    契约：
      - 依次把每个 input 的所有页加进同一个 PdfWriter；
      - 写到 output_path（用 ``with open(output_path, "wb") as f: writer.write(f)``）；
      - input_paths 为空时写出空 PDF（0 页）；
      - 返回 output_path。

    提示：
      - writer.append(PdfReader(p)) 一行搞定一个文件全部页（保书签）；
        也可 for page in PdfReader(p).pages: writer.add_page(page)。
      - 测试断言：输出页数 == 各输入页数之和。
    """
    writer = PdfWriter()
    for p in input_paths:
        writer.append(PdfReader(p))
    with open(output_path, "wb") as f:
        writer.write(f)
    return output_path
    raise NotImplementedError("TODO(Layer a): 实现 merge_pdfs")


def split_pdf(input_path: str, output_dir: str) -> list[str]:
    """[学习·Layer (a) TODO] 把 PDF 每页拆成一个单独 PDF，返回有序路径列表。

    契约：
      - 第 i 页（从 1 计）写出为 output_dir/page_{i}.pdf；
      - 返回按页序排列的输出路径列表。

    提示：
      - 遍历 PdfReader(input_path).pages（或 enumerate(..., start=1)）；
      - 每页 new 一个 PdfWriter、add_page(page)、写出。
    """
    reader = PdfReader(input_path)
    paths = []
    for i, page in enumerate(reader.pages, start=1):
        # 每页独立 writer，否则页会累积（page_2 会含 page_1，page_3 含 1+2）。
        writer = PdfWriter()
        writer.add_page(page)
        out = f"{output_dir}/page_{i}.pdf"
        with open(out, "wb") as f:
            writer.write(f)
        paths.append(out)
    return paths
    raise NotImplementedError("TODO(Layer a): 实现 split_pdf")


def rotate_pdf(input_path: str, output_path: str, angle: int = 90) -> str:
    """[学习·Layer (a) TODO] 旋转所有页，写出后返回 output_path。

    契约：
      - angle 只接受 {90, 180, 270}；其它值 raise UnsupportedActionError 或 ValueError；
      - 把每页旋转后加到新 PdfWriter，写出 output_path；
      - 返回 output_path。

    提示：page.rotate(angle) 原地修改并返回 page；把它 add_page 到新 writer。
    """
    if angle!=90 and angle!=180 and angle!=270:
        raise UnsupportedActionError("请使用正确的旋转角度(90,180,270)")
    reader = PdfReader(input_path)
    writer = PdfWriter()
    for page in reader.pages:
        page.rotate(angle)
        writer.add_page(page)
    with open(output_path, "wb") as f:
        writer.write(f)
    return output_path     
    raise NotImplementedError("TODO(Layer a): 实现 rotate_pdf")


def extract_pdf_pages(input_path: str, page_indices: list[int], output_path: str) -> str:
    """[学习·Layer (a) TODO] 按 0-based 页码抽取页面合并为新 PDF，返回 output_path。

    契约：
      - page_indices 是 0-based 索引列表；
      - 把这些页按给定顺序加进新 PdfWriter，写出 output_path；
      - 返回 output_path。

    提示：pages = PdfReader(input_path).pages; writer.add_page(pages[i])。
    """
    writer = PdfWriter()
    for i in page_indices:
        writer.add_page(PdfReader(input_path).pages[i])
    with open(output_path, "wb") as f:
        writer.write(f)
    return output_path  
    raise NotImplementedError("TODO(Layer a): 实现 extract_pdf_pages")


def pdf_to_images(
    input_path: str, output_dir: str, zoom: float = 2.0, fmt: str = "png"
) -> list[str]:
    """[学习·Layer (a) TODO] 逐页渲染为图片，返回有序路径列表。

    契约：
      - 第 i 页（从 1 计）渲染为 output_dir/page_{i}.{fmt}；
      - zoom 控制清晰度（Matrix(zoom, zoom)）；
      - 返回按页序排列的输出路径列表。

    提示：
      - import fitz; doc = fitz.open(input_path);
      - page.get_pixmap(matrix=fitz.Matrix(zoom, zoom)) → pix.save(path)；
      - pix.save 按扩展名定格式，所以 path 后缀用 .{fmt}。
    """
    doc = fitz.open(input_path)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        pix.save(f"{output_dir}/page_{i+1}.{fmt}")
    return [f"{output_dir}/page_{i+1}.{fmt}" for i in range(len(doc))]
    raise NotImplementedError("TODO(Layer a): 实现 pdf_to_images")


def extract_pdf_text(input_path: str) -> str:
    """[学习·Layer (a) TODO] 提取全文，页间用 \\n 分隔，返回 str。

    契约：把每页 page.get_text() 用 "\\n" 拼接；无文本层返回空串不报错。

    提示：fitz.open(input_path) 后遍历 doc，page.get_text()。也可用 pypdf 的
    page.extract_text()，但 fitz 对中文/排版更稳，推荐 fitz。
    """
    doc = fitz.open(input_path)
    text = ""
    for page in doc:
        text += page.get_text()
        text += "\n"
    return text
    raise NotImplementedError("TODO(Layer a): 实现 extract_pdf_text")


# ===== 扩展功能：加密 / 解密 / 水印 / 页码 =====


def encrypt_pdf(
    input_path: str,
    output_path: str,
    user_password: str,
    owner_password: str | None = None,
) -> str:
    """给 PDF 加密码，返回 output_path。

    契约：把每页复制进新 writer，``writer.encrypt(user_password, owner_password)``
    （owner_password 留空则与 user_password 相同），写出到 output_path。
    """
    reader = PdfReader(input_path)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(
        user_password=user_password,
        owner_password=owner_password or user_password,
    )
    with open(output_path, "wb") as f:
        writer.write(f)
    return output_path


def decrypt_pdf(input_path: str, output_path: str, password: str) -> str:
    """去掉 PDF 密码（用给定密码解密后重写），返回 output_path。

    契约：``reader.is_encrypted`` 为真时先 ``reader.decrypt(password)``，
    再把解密后的页写进不带加密的新 writer。
    """
    reader = PdfReader(input_path)
    if reader.is_encrypted:
        reader.decrypt(password)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    with open(output_path, "wb") as f:
        writer.write(f)
    return output_path


def add_watermark(
    input_path: str,
    output_path: str,
    text: str,
    *,
    font_size: float = 48,
    opacity: float = 0.25,
    angle: int = 45,
) -> str:
    """给每页叠加半透明文字水印（默认斜 45°），返回 output_path。

    用 fitz 的 ``fill_opacity`` + ``morph``（绕页面中心旋转）实现斜体水印；
    ``angle`` 取 0 时水平居中。文字是真实文本，可被 ``get_text()`` 提取。
    """
    doc = fitz.open(input_path)
    for page in doc:
        rect = page.rect
        width = fitz.get_text_length(text, fontsize=font_size)
        start = fitz.Point((rect.width - width) / 2, rect.height / 2)
        pivot = fitz.Point(rect.width / 2, rect.height / 2)
        # rotate 只接受 90 的倍数，任意角度用 morph(绕 pivot 的旋转矩阵)
        page.insert_text(
            start,
            text,
            fontsize=font_size,
            color=(0.6, 0.6, 0.6),
            fill_opacity=opacity,
            morph=(pivot, fitz.Matrix(angle)) if angle else None,
        )
    doc.save(output_path)
    doc.close()
    return output_path


def add_page_numbers(
    input_path: str,
    output_path: str,
    *,
    start: int = 1,
    font_size: float = 11,
    margin: float = 28,
) -> str:
    """在每页底部居中加页码（从 start 起），返回 output_path。"""
    doc = fitz.open(input_path)
    for i, page in enumerate(doc, start=start):
        rect = page.rect
        label = str(i)
        width = fitz.get_text_length(label, fontsize=font_size)
        page.insert_text(
            fitz.Point((rect.width - width) / 2, rect.height - margin / 2),
            label,
            fontsize=font_size,
            color=(0.2, 0.2, 0.2),
        )
    doc.save(output_path)
    doc.close()
    return output_path


def _xml_safe(text: str) -> str:
    """去掉 XML 1.0 不允许的字符（控制字符/半区代理），否则 python-docx 会抛 ValueError。"""
    return "".join(
        ch
        for ch in text
        if ch in "\t\n\r"
        or 0x20 <= ord(ch) <= 0xD7FF
        or 0xE000 <= ord(ch) <= 0xFFFD
        or 0x10000 <= ord(ch) <= 0x10FFFF
    )


def pdf_to_docx(input_path: str, output_path: str, *, include_images: bool = True) -> str:
    """把 PDF 转成 Word（.docx），返回 output_path。

    逐页提取文本块，按行生成段落，尽量保留字号/加粗/斜体；页间插入分页符；
    ``include_images`` 为真时把每页图片附在页尾。基于文本层的转换，复杂版式/表格/多栏
    可能不完美；纯扫描件（无文本层）需先走 OCR。
    """
    src = fitz.open(input_path)
    doc = Document()
    for pno, page in enumerate(src, start=1):
        data = page.get_text("dict")
        for block in data.get("blocks", []):
            if block.get("type") != 0:  # 0 = 文本块
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                # PDF 文本层可能夹带控制字符，先清洗成 XML 允许的文本再写 docx。
                texts = [_xml_safe(s.get("text", "")) for s in spans]
                if not "".join(texts).strip():
                    continue
                para = doc.add_paragraph()
                for span, text in zip(spans, texts):
                    if not text:
                        continue
                    run = para.add_run(text)
                    size = span.get("size")
                    if size:
                        run.font.size = Pt(round(float(size), 1))
                    flags = span.get("flags", 0)
                    run.bold = bool(flags & 16)    # bit4 = 粗体
                    run.italic = bool(flags & 2)   # bit1 = 斜体
        if include_images:
            for img in page.get_images(full=True):
                try:
                    info = src.extract_image(img[0])
                except Exception as exc:  # noqa: BLE001 — 坏图跳过，不中断整篇转换
                    logger.warning("PDF 转 Word: 提取图片失败，跳过: {}", exc)
                    continue
                blob = info.get("image")
                if not blob:
                    continue
                try:
                    doc.add_picture(BytesIO(blob), width=Inches(6))
                except Exception as exc:  # noqa: BLE001 — 个别格式不支持则跳过
                    logger.warning("PDF 转 Word: 插入图片失败，跳过: {}", exc)
                    continue
        if pno < src.page_count:
            doc.add_page_break()
    doc.save(output_path)
    src.close()
    return output_path


# ===== 已写好：PdfProcessor 分派胶水（学生不要改） =====


class PdfProcessor(FileProcessor):
    supported_extensions = {".pdf"}

    def run(self, file_path: str, action: str, options: dict) -> Any:
        match action:
            case "merge":
                return merge_pdfs(options["inputs"], options["output"])
            case "split":
                return split_pdf(file_path, options["output_dir"])
            case "rotate":
                return rotate_pdf(file_path, options["output"], options.get("angle", 90))
            case "extract_pages":
                return extract_pdf_pages(file_path, options["pages"], options["output"])
            case "to_images":
                return pdf_to_images(
                    file_path,
                    options["output_dir"],
                    options.get("zoom", 2.0),
                    options.get("fmt", "png"),
                )
            case "extract_text":
                return extract_pdf_text(file_path)
            case "encrypt":
                return encrypt_pdf(
                    file_path,
                    options["output"],
                    options["password"],
                    options.get("owner_password"),
                )
            case "decrypt":
                return decrypt_pdf(file_path, options["output"], options["password"])
            case "watermark":
                return add_watermark(
                    file_path,
                    options["output"],
                    options.get("text") or "CONFIDENTIAL",
                    font_size=options.get("font_size", 48),
                    opacity=options.get("opacity", 0.25),
                )
            case "page_numbers":
                return add_page_numbers(
                    file_path,
                    options["output"],
                    start=options.get("start", 1),
                )
            case "to_word":
                return pdf_to_docx(
                    file_path,
                    options["output"],
                    include_images=options.get("include_images", True),
                )
            case _:
                raise UnsupportedActionError(f"PdfProcessor 不支持操作: {action}")


__all__ = [
    "PdfProcessor",
    "add_page_numbers",
    "add_watermark",
    "decrypt_pdf",
    "encrypt_pdf",
    "extract_pdf_pages",
    "extract_pdf_text",
    "merge_pdfs",
    "pdf_to_docx",
    "pdf_to_images",
    "rotate_pdf",
    "split_pdf",
]
