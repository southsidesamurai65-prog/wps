"""DOCX 处理器。

Layer (a) TODO：extract_docx_text / replace_docx_text 函数体。
DocxProcessor 类已写好当分派胶水。

库 API 要点（python-docx 1.2）：
- Document(path) 读；document.paragraphs 遍历段。
- ⚠️ paragraph.text = ... 会把整段文本塞进第一个 run 并清空其余 run（丢 run 样式）。
  因此 replace 必须遍历 run 做替换（run.text = run.text.replace(old, new)），
  否则 tests/test_docx_processor.py 的「保样式档」会因 run 数变少而红——
  这正是 spec 用来抓「粗暴 paragraph.text= 替换」常见错误的设计。
- 标题检测用 paragraph.style.name（"Title"/"Heading 1"），别靠缩进。
"""

from __future__ import annotations

from typing import Any

from wps_tool.core.errors import UnsupportedActionError
from wps_tool.processors.base import FileProcessor

from docx import Document


def extract_docx_text(input_path: str) -> str:
    """[学习·Layer (a) TODO] 提取正文文本，段间用 \\n 连接，返回 str。

    契约：取 document.paragraphs 中非空段，p.text，用 "\\n".join。

    提示：from docx import Document; Document(input_path).paragraphs。
    """
    doc = Document(input_path)
    text = "\n".join([p.text for p in doc.paragraphs if p.text])
    return text
    raise NotImplementedError("TODO(Layer a): 实现 extract_docx_text")


def replace_docx_text(input_path: str, output_path: str, mapping: dict[str, str]) -> str:
    """[学习·Layer (a) TODO] 按 mapping 批量替换文本，保存到 output_path，返回它。

    契约：
      - 对每个 paragraph，遍历它的 runs，对每个 run 做
        run.text = run.text.replace(old, new)（old/new 来自 mapping）；
      - ⚠️ 不要用 paragraph.text = paragraph.text.replace(...)——会丢 run 样式，
        test_docx_processor.py 的保样式档会因 run 数变少而失败；
      - 保存 document.save(output_path)；返回 output_path。

    提示：
      - 一个 token 可能跨多个 run（python-docx 把同一文本拆进多个 run），
        M1–M3 不要求处理跨 run token，只要 token 完整落在一个 run 内即可替换；
      - 测试夹具把 {{COMPANY}} 放在单 run 内，故逐 run 替换能命中。
    """
    doc = Document(input_path)
    for p in doc.paragraphs:
        for r in p.runs: #run是文本，比段落还小
            for k, v in mapping.items():
                r.text = r.text.replace(k, v)
    doc.save(output_path)
    return output_path
    raise NotImplementedError("TODO(Layer a): 实现 replace_docx_text")


class DocxProcessor(FileProcessor):
    supported_extensions = {".docx"}

    def run(self, file_path: str, action: str, options: dict) -> Any:
        match action:
            case "extract_text":
                return extract_docx_text(file_path)
            case "replace":
                return replace_docx_text(file_path, options["output"], options["mapping"])
            case _:
                raise UnsupportedActionError(f"DocxProcessor 不支持操作: {action}")


__all__ = ["DocxProcessor", "extract_docx_text", "replace_docx_text"]
