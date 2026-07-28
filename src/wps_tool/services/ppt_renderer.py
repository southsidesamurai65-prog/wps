"""PPT 美化本地渲染器（纯函数，无网络）—— shape 级。

把 LLM 返回的「重设计 spec JSON」（shape 级，含每个 shape 的坐标/尺寸/类型/样式）
用 python-pptx 在本地重建 deck。pptx 字节不出本机——spec 只含文本/版式/配色/坐标，
渲染全在本地完成，可被单测直接驱动。图片由调用方从原 pptx 本地抽图后按 ``image_id``
传入（blob 留在本机内存，不进 spec、不发网络）——隐私：LLM 看不到图片内容。

spec 协议（LLM 返回）::

    {
      "slide_size": {"width": 13.333, "height": 7.5},
      "theme": {"accent": "#1F4E79", "bg": "#FFFFFF"},
      "slides": [
        {"bg": "#FFFFFF", "shapes": [
          {"type": "rect", "left": 0, "top": 0, "width": 13.333, "height": 1.4,
           "fill": "#1F4E79", "line": null, "text": "标题", "font_size": 40,
           "font_color": "#FFFFFF", "bold": true, "align": "left",
           "valign": "middle", "margin_left": 0.5},
          {"type": "image", "image_id": "img_1", "left": 7.6, "top": 2,
           "width": 5, "height": 4.5},
          {"type": "textbox", "left": 0.6, "top": 2, "width": 6.5, "height": 4.5,
           "valign": "top", "paragraphs": [
             {"text": "收入 +18%", "bullet": true, "font_size": 22, "color": "#222222"}
           ]}
        ]}
      ]
    }

约束：slides 数应 == 源页数；每 shape 有 ``type``+``left/top/width/height``（英寸）；
image shape 用 ``image_id`` 引用调用方传入的 ``images`` 字典。未知 type 当 textbox 兜底。
"""

from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from wps_tool.utils.logging import logger

#: 默认 accent 色（spec 缺 theme 时兜底，深商务蓝）。
_DEFAULT_ACCENT = "#1F4E79"
#: 默认背景色（spec 缺 theme 时兜底，纯白）。
_DEFAULT_BG = "#FFFFFF"
#: 正文文字色（深灰，不刺眼）。
_TEXT_DARK = "#222222"
#: EMU per inch（python-pptx 内部长度单位：914400 EMU = 1 英寸）。
_EMU_IN = 914400

#: shape type → python-pptx 自选图形枚举。
_SHAPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
}

_ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}

_VALIGN_MAP = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def _hex(color: str | None, fallback: str) -> RGBColor:
    """``"#1F4E79"`` / ``"1F4E79"`` / ``None``（用 fallback）→ ``RGBColor``。"""
    c = color or fallback
    return RGBColor.from_string(c.lstrip("#"))


def _hex_opt(color: str | None, fallback: RGBColor) -> RGBColor:
    """``color`` 非空则取，否则用已解析的 ``fallback``。"""
    if not color:
        return fallback
    return RGBColor.from_string(color.lstrip("#"))


def _num(shape_spec: dict, key: str, default: float) -> float:
    """取 shape_spec 里的数值，缺/非数时回退 ``default``。"""
    v = shape_spec.get(key, default)
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _set_bg(slide, color: RGBColor) -> None:
    """整页背景填 ``color``（覆盖 layout/master 的背景）。"""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _style_run(run, *, size: int, color: RGBColor, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _apply_valign(shape_spec: dict, text_frame) -> None:
    valign = shape_spec.get("valign")
    if valign in _VALIGN_MAP:
        text_frame.vertical_anchor = _VALIGN_MAP[valign]


def _apply_single_text(shape_spec: dict, text_frame, *, default_size: int,
                       default_color: RGBColor) -> None:
    """把 shape_spec 的 ``text`` + 对齐/字号/色/粗体套到单段 text_frame。"""
    text_frame.word_wrap = True
    _apply_valign(shape_spec, text_frame)
    ml = shape_spec.get("margin_left")
    if ml is not None:
        text_frame.margin_left = Inches(float(ml))
    p = text_frame.paragraphs[0]
    p.text = shape_spec.get("text") or ""
    align = shape_spec.get("align")
    if align in _ALIGN_MAP:
        p.alignment = _ALIGN_MAP[align]
    size = int(shape_spec.get("font_size") or default_size)
    color = _hex_opt(shape_spec.get("font_color"), default_color)
    bold = bool(shape_spec.get("bold"))
    for run in p.runs:
        _style_run(run, size=size, color=color, bold=bold)


def _add_autoshape(slide, shape_spec: dict, accent: RGBColor, white: RGBColor) -> None:
    """rect / rounded_rect / oval：自选图形 + 填充 + 描边 + 可选单段文字。"""
    mso = _SHAPE_MAP.get(shape_spec.get("type", "rect"), MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(
        mso,
        Inches(_num(shape_spec, "left", 0)),
        Inches(_num(shape_spec, "top", 0)),
        Inches(_num(shape_spec, "width", 4)),
        Inches(_num(shape_spec, "height", 1)),
    )
    fill = shape.fill
    fill.solid()
    fill_hex = shape_spec.get("fill")
    fill.fore_color.rgb = RGBColor.from_string(fill_hex.lstrip("#")) if fill_hex else accent
    line = shape_spec.get("line")
    if line is None:
        shape.line.fill.background()  # 无描边
    else:
        shape.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
    if shape_spec.get("text"):
        _apply_single_text(shape_spec, shape.text_frame,
                           default_size=24, default_color=white)


def _add_textbox(slide, shape_spec: dict, dark: RGBColor) -> None:
    """textbox：多段文字（每段可选 bullet、字号、色），可设纵向对齐。"""
    tb = slide.shapes.add_textbox(
        Inches(_num(shape_spec, "left", 0)),
        Inches(_num(shape_spec, "top", 0)),
        Inches(_num(shape_spec, "width", 6)),
        Inches(_num(shape_spec, "height", 2)),
    )
    tf = tb.text_frame
    tf.word_wrap = True
    _apply_valign(shape_spec, tf)
    paragraphs = shape_spec.get("paragraphs") or []
    if not paragraphs and shape_spec.get("text"):
        paragraphs = [{"text": shape_spec["text"]}]
    for i, para in enumerate(paragraphs):
        if not isinstance(para, dict):
            continue
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text = para.get("text", "") or ""
        if para.get("bullet"):
            text = f"• {text}"
        p.text = text
        size = int(para.get("font_size") or 18)
        color = _hex_opt(para.get("color"), dark)
        for run in p.runs:
            _style_run(run, size=size, color=color)


def _add_image(slide, shape_spec: dict, images: dict) -> None:
    """image：按 image_id 从本地 images 字典取 blob 插入（blob 不进 spec、不发网络）。"""
    image_id = shape_spec.get("image_id")
    blob = images.get(image_id) if (images and image_id) else None
    if blob is None:
        logger.warning("PPT 渲染: 图片缺失，跳过: image_id={}", image_id)
        return
    width = _num(shape_spec, "width", 0) or None
    height = _num(shape_spec, "height", 0) or None
    slide.shapes.add_picture(
        BytesIO(blob),
        Inches(_num(shape_spec, "left", 0)),
        Inches(_num(shape_spec, "top", 0)),
        width=Inches(width) if width else None,
        height=Inches(height) if height else None,
    )


def render_beautified_deck(
    spec: dict,
    output_path: str,
    style: str = "business",
    *,
    slide_size: dict | None = None,
    images: dict | None = None,
) -> str:
    """把 LLM 返回的 shape 级 ``spec`` 渲染成 deck，写到 ``output_path``，返回它。

    纯本地、无网络——可被单测直接驱动。
      - ``slide_size``：``{"width", "height"}``（英寸），设新 deck 画布同源尺寸。
      - ``images``：``{image_id: bytes}``，image shape 按 id 取 blob 本地插入。
    ``style`` 当前只影响语义，预留以便后续按风格切配色。
    """
    theme = (spec.get("theme") or {}) if isinstance(spec, dict) else {}
    accent = _hex(theme.get("accent"), _DEFAULT_ACCENT)
    bg_fallback = theme.get("bg") or _DEFAULT_BG
    white = _hex("FFFFFF", "FFFFFF")
    dark = _hex(_TEXT_DARK, _TEXT_DARK)
    images = images or {}
    slide_size = slide_size or spec.get("slide_size")

    prs = Presentation()
    if slide_size:
        prs.slide_width = Inches(float(slide_size.get("width", 10)))
        prs.slide_height = Inches(float(slide_size.get("height", 7.5)))
    for slide_spec in spec.get("slides", []) or []:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        _set_bg(slide, _hex(slide_spec.get("bg"), bg_fallback))
        for shape_spec in slide_spec.get("shapes", []) or []:
            if not isinstance(shape_spec, dict):
                continue
            stype = (shape_spec.get("type") or "textbox").strip().lower()
            if stype in _SHAPE_MAP:
                _add_autoshape(slide, shape_spec, accent, white)
            elif stype == "image":
                _add_image(slide, shape_spec, images)
            else:  # textbox / 未知 type 兜底
                _add_textbox(slide, shape_spec, dark)

    prs.save(output_path)
    return output_path


__all__ = ["render_beautified_deck"]
