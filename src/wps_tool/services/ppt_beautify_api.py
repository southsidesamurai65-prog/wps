"""PPT 美化：LLM + 本地渲染（pptx 字节不出本机）—— shape 级。

流程：本地解析 pptx 的每页 shape 清单（文本/图片占位）→ 把清单（**不含图片字节**，
只含每页文本 + 图片的 image_id/位置/尺寸）发给 LLM（OpenAI Chat Completions，
``response_format=json_object`` 强制 JSON）→ 拿回「shape 级重设计 spec」→
用 python-pptx 在本地重建 deck（``render_beautified_deck``），图片从原 pptx 本地抽图、
按 image_id 插入新位置。

隐私：原 pptx 的二进制本体绝不上传，图片字节也只留本机内存（不进请求体）。
请求体里只有每页文本 + 图片占位信息（断言**不含** pptx 的 zip 头 ``PK\\x03\\x04``，
也**不含** PNG 头 ``\\x89PNG``，见 ``tests/test_ppt_beautify_api.py``）。

关键设计（httpx 0.28，测试离线）：客户端必须持有「可注入 transport 的 httpx.Client」，
而不是用模块级 httpx.post（瞬时 Client 无法注入 transport，无法离线测试）。
测试用 ``httpx.MockTransport(handler)`` 注入，handler 签名 ``(req) -> Response``，
并在 handler 里断言请求体**不含** PK/PNG 头、URL/鉴权头正确。
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import httpx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from wps_tool.core.errors import ApiUnavailableError
from wps_tool.services.ppt_renderer import render_beautified_deck
from wps_tool.utils.logging import logger

#: pptx 文件的 zip 头字节（PK\\x03\\x04）。隐私断言用：LLM 请求体**不应**含它。
PPTX_MAGIC = b"PK\x03\x04"

#: PNG 文件头字节（\\x89PNG）。隐私断言用：图片字节不出本机，LLM 请求体**不应**含它。
PNG_MAGIC = b"\x89PNG"

#: OpenAI Chat Completions 端点（相对 base_url）。
_CHAT_COMPLETIONS_PATH = "/v1/chat/completions"

#: 当前实现的 LLM 供应商；其它 provider 会抛清晰错。
_SUPPORTED_PROVIDER = "openai"

#: EMU per inch（python-pptx 内部长度单位：914400 EMU = 1 英寸）。
_EMU_IN = 914400

_SYSTEM_PROMPT = (
    "你是资深 PPT 设计师。把用户给的每页 shape 清单重设计成一套干净、一致的商务 deck。"
    "严格输出 JSON，schema：\n"
    '{"slide_size": {"width": 13.333, "height": 7.5}, '
    '"theme": {"accent": "#RRGGBB", "bg": "#RRGGBB"}, "slides": [\n'
    '  {"bg": "#FFFFFF", "shapes": [\n'
    '    {"type": "rect", "left": 0, "top": 0, "width": 13.333, "height": 1.4, '
    '"fill": "#1F4E79", "line": null, "text": "标题", "font_size": 40, '
    '"font_color": "#FFFFFF", "bold": true, "align": "left", "valign": "middle", '
    '"margin_left": 0.5},\n'
    '    {"type": "image", "image_id": "img_1", "left": 7.6, "top": 2, '
    '"width": 5, "height": 4.5},\n'
    '    {"type": "textbox", "left": 0.6, "top": 2, "width": 6.5, "height": 4.5, '
    '"valign": "top", "paragraphs": [{"text": "要点", "bullet": true, '
    '"font_size": 22, "color": "#222222"}]}\n'
    "  ]}\n"
    "]}\n"
    "约束：slides 数量必须 == 输入页数（1:1，逐页对应，顺序一致）；"
    "type 取 rect / rounded_rect / oval / textbox / image；坐标/尺寸单位英寸，"
    "不要超出画布；image 必须用输入里出现的 image_id 引用，不要新造 id；"
    "align ∈ left/center/right，valign ∈ top/middle/bottom，line 为 null 表示无线；"
    "accent 用饱和度适中的深色，bg 用浅色。"
)


def _extract_shape_manifest(input_path: str):
    """本地解析原 deck：每页 shape 清单（文本/图片占位）+ 图片 blob 字典 + 画布尺寸。

    返回 ``(pages, images, slide_size)``：
      - ``pages``：``[{"slide": i, "shapes": [{"kind": "text"|"image",
        "left","top","width","height" (英寸), "text"?, "image_id"?}]}]``
      - ``images``：``{"img_1": bytes, ...}``——图片 blob，**留本机内存，不进请求体**
      - ``slide_size``：``{"width", "height"}``（英寸）

    图片按出现顺序分配 ``image_id=f"img_{n}"``；取 blob 失败的图跳过并告警。
    """
    prs = Presentation(input_path)
    slide_size = {
        "width": prs.slide_width / _EMU_IN,
        "height": prs.slide_height / _EMU_IN,
    }
    pages: list[dict] = []
    images: dict[str, bytes] = {}
    img_n = 0
    for i, slide in enumerate(prs.slides, start=1):
        shapes: list[dict] = []
        for shape in slide.shapes:
            left = shape.left / _EMU_IN if shape.left is not None else 0.0
            top = shape.top / _EMU_IN if shape.top is not None else 0.0
            width = shape.width / _EMU_IN if shape.width is not None else 0.0
            height = shape.height / _EMU_IN if shape.height is not None else 0.0
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_n += 1
                image_id = f"img_{img_n}"
                try:
                    images[image_id] = shape.image.blob
                except (AttributeError, KeyError, ValueError):  # 取 blob 失败：跳过该图
                    logger.warning(
                        "PPT 解析: 取图片 blob 失败，跳过: slide={} image_id={}",
                        i,
                        image_id,
                    )
                    continue
                shapes.append(
                    {"kind": "image", "image_id": image_id,
                     "left": left, "top": top, "width": width, "height": height}
                )
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    shapes.append(
                        {"kind": "text", "text": text,
                         "left": left, "top": top, "width": width, "height": height}
                    )
        pages.append({"slide": i, "shapes": shapes})
    return pages, images, slide_size


def _build_user_prompt(pages: list[dict], slide_size: dict, style: str) -> str:
    """把本地解析的每页 shape 清单拼成给 LLM 的 user 消息。

    图片只给 ``image_id`` + 位置/尺寸（内容 LLM 看不到），要求 LLM 用相同 id 引用、
    重新摆位；文本给原文，要求润色保留。
    """
    lines = [
        f"画布尺寸：{slide_size['width']:.2f} × {slide_size['height']:.2f} 英寸。",
        f"风格：{style}。",
        (
            f"共 {len(pages)} 页。每页 shape 清单如下（坐标英寸；图片内容你看不到，"
            "只给位置/尺寸，用 image_id 引用，必须保留并在新 spec 里用相同 image_id 引用）："
        ),
    ]
    for page in pages:
        lines.append(f"第{page['slide']}页：")
        for sh in page["shapes"]:
            geo = (
                f"位置({sh['left']:.2f},{sh['top']:.2f}) "
                f"尺寸 {sh['width']:.2f}×{sh['height']:.2f}"
            )
            if sh["kind"] == "image":
                lines.append(
                    f"  - [图片] image_id={sh['image_id']} {geo}（内容不可见，重新摆位即可）"
                )
            else:
                lines.append(f"  - [文本] {sh['text']} {geo}")
    lines.append(
        "请把每页重设计为 shape 级 spec：保留所有文本内容（可润色更精炼）、"
        "保留所有图片（用相同 image_id 引用、放合理新位置）、坐标对齐 0.5 英寸网格、"
        "标题层级清晰、accent 只做强调、图片不压文字。只输出 JSON。"
    )
    return "\n".join(lines)


class PptBeautifyClient:
    """PPT 美化客户端：本地解析 → LLM 拿 shape 级 spec → 本地渲染重建 deck。"""

    def __init__(
        self,
        api_key: str,
        model: str = "gpt-4o",
        *,
        provider: str = _SUPPORTED_PROVIDER,
        base_url: str = "https://api.openai.com",
        transport: httpx.BaseTransport | None = None,
        timeout: float = 120,
    ) -> None:
        if provider and provider.strip().lower() != _SUPPORTED_PROVIDER:
            raise ApiUnavailableError(
                f"PPT 美化当前只支持 provider=openai，得到：{provider!r}"
            )
        self.api_key = api_key
        self.model = model
        self.provider = _SUPPORTED_PROVIDER
        # 持有可注入 Client：测试传 transport=httpx.MockTransport(...) 即可离线。
        self._client = httpx.Client(
            base_url=base_url.rstrip("/"),
            transport=transport,
            timeout=timeout,
        )
        logger.debug(
            "PPT 美化客户端初始化: provider={} model={} base_url={} timeout={}",
            self.provider,
            self.model,
            str(self._client.base_url).rstrip("/"),
            timeout,
        )

    @property
    def client(self) -> httpx.Client:
        """暴露底层 Client，供测试读取（如断言发出的请求）。"""
        return self._client

    def _auth_headers(self) -> dict[str, str]:
        return {"Authorization": f"Bearer {self.api_key}"}

    def close(self) -> None:
        self._client.close()

    def _call_llm(self, user_prompt: str) -> dict[str, Any]:
        """POST OpenAI Chat Completions，解析返回的 spec JSON。"""
        body = {
            "model": self.model,
            "messages": [
                {"role": "system", "content": _SYSTEM_PROMPT},
                {"role": "user", "content": user_prompt},
            ],
            "response_format": {"type": "json_object"},
        }
        logger.info(
            "PPT 美化 LLM 请求准备: method=POST base_url={} endpoint={} model={} prompt_chars={}",
            str(self._client.base_url).rstrip("/"),
            _CHAT_COMPLETIONS_PATH,
            self.model,
            len(user_prompt),
        )
        try:
            response = self._client.post(
                _CHAT_COMPLETIONS_PATH,
                json=body,
                headers=self._auth_headers(),
            )
        except httpx.HTTPError:
            logger.exception(
                "PPT 美化 LLM 请求失败: endpoint={}", _CHAT_COMPLETIONS_PATH
            )
            raise
        logger.info(
            "PPT 美化 LLM 收到响应: status_code={} content_type={}",
            response.status_code,
            response.headers.get("content-type", ""),
        )
        try:
            response.raise_for_status()
        except httpx.HTTPStatusError:
            logger.error(
                "PPT 美化 LLM HTTP 错误: status_code={} body_preview={}",
                response.status_code,
                response.text[:500],
            )
            raise
        try:
            data = response.json()
            content = data["choices"][0]["message"]["content"]
            spec = json.loads(content)
        except (KeyError, IndexError, TypeError, json.JSONDecodeError):
            logger.exception(
                "PPT 美化 LLM 响应解析失败: status_code={} body_chars={}",
                response.status_code,
                len(response.text),
            )
            raise
        if not isinstance(spec, dict) or "slides" not in spec:
            logger.error(
                "PPT 美化 LLM spec 不合法: spec_type={} has_slides={}",
                type(spec).__name__,
                isinstance(spec, dict) and "slides" in spec,
            )
            raise ApiUnavailableError(f"LLM 返回的 spec 不合法：{content!r}")
        logger.info(
            "PPT 美化 LLM spec 解析完成: slides={}",
            len(spec.get("slides", [])),
        )
        return spec

    def beautify_file(
        self, input_path: str, output_path: str, style: str = "business"
    ) -> str:
        """美化一个 pptx：本地解析 shape 清单 → 文本+图片占位发 OpenAI → 本地渲染重建 deck。

        契约：
          - 本地解析：``_extract_shape_manifest`` 拿每页 shape 清单 + 图片 blob（本机）+ 画布尺寸；
          - 把每页文本 + 图片占位（``image_id``/位置/尺寸，**不含图片字节**）发 OpenAI
            ``/v1/chat/completions``，``response_format=json_object``，headers ``Authorization: Bearer {key}``；
          - 拿回 shape 级 spec → ``render_beautified_deck(spec, output, style, images=images)``；
          - 返回 ``output_path``。
          - api_key 为空 / provider 非 openai 时 raise ``ApiUnavailableError``。
          - 请求体**不含** pptx 二进制也**不含**图片字节（隐私：只发文本+图片占位）。
        """
        if not self.api_key:
            logger.error("PPT 美化中止: api_key 为空")
            raise ApiUnavailableError("api_key 不能为空")
        logger.info(
            "PPT 美化开始: input_name={} output={} style={}",
            Path(input_path).name,
            output_path,
            style,
        )
        pages, images, slide_size = _extract_shape_manifest(input_path)
        logger.info(
            "PPT 美化本地解析完成: input_name={} pages={} images={}",
            Path(input_path).name,
            len(pages),
            len(images),
        )
        user_prompt = _build_user_prompt(pages, slide_size, style)
        logger.debug("PPT 美化 prompt 构建完成: prompt_chars={}", len(user_prompt))
        spec = self._call_llm(user_prompt)
        # 确保 spec 带画布尺寸（渲染器要用）；LLM 没给就用源的。
        if "slide_size" not in spec:
            spec = {**spec, "slide_size": slide_size}
        # spec 页数应 == 源页数；缺失或对不齐时，按源页数补齐/截断，保证产物可打开。
        src_pages = len(pages)
        slides = spec.get("slides", [])
        if len(slides) != src_pages:
            logger.warning(
                "PPT 美化 spec 页数不匹配: src_pages={} spec_pages={} 将自动对齐",
                src_pages,
                len(slides),
            )
            spec = {**spec, "slides": _align_pages(slides, src_pages, pages, slide_size)}
        result = render_beautified_deck(spec, output_path, style, images=images)
        logger.info(
            "PPT 美化本地渲染完成: output={} slides={}",
            result,
            len(spec.get("slides", [])),
        )
        return result


def _align_pages(
    slides: list[dict], n: int, pages: list[dict], slide_size: dict
) -> list[dict]:
    """spec 页数与源页数对不齐时的兜底：截断到 n，或用源 shape 清单生成最小 spec slide。

    补的页：accent 标题 rect（取源页第一个文本）+ 文本 textbox（其余文本作 bullet）
    + 原图图片（用相同 image_id 搬到原位置）。保证产物页数对、文本/图片不丢。
    """
    aligned = list(slides[:n])
    for i in range(len(aligned), n):
        page = pages[i] if i < len(pages) else {"shapes": []}
        page_shapes = page.get("shapes", [])
        texts = [s for s in page_shapes if s.get("kind") == "text"]
        img_shapes = [s for s in page_shapes if s.get("kind") == "image"]
        new_shapes: list[dict] = []
        if texts:
            new_shapes.append(
                {"type": "rect", "left": 0, "top": 0,
                 "width": slide_size["width"], "height": 1.4,
                 "fill": "#1F4E79", "text": texts[0]["text"],
                 "font_color": "#FFFFFF", "bold": True,
                 "align": "left", "valign": "middle", "margin_left": 0.5}
            )
            bullets = [t["text"] for t in texts[1:]]
        else:
            bullets = [f"第{i + 1}页"]
        if bullets:
            new_shapes.append(
                {"type": "textbox", "left": 0.6, "top": 2,
                 "width": slide_size["width"] - 1.2,
                 "height": slide_size["height"] - 2.5, "valign": "top",
                 "paragraphs": [{"text": b, "bullet": True} for b in bullets]}
            )
        for s in img_shapes:
            new_shapes.append(
                {"type": "image", "image_id": s["image_id"],
                 "left": s["left"], "top": s["top"],
                 "width": s["width"], "height": s["height"]}
            )
        aligned.append({"shapes": new_shapes})
    return aligned


__all__ = ["PNG_MAGIC", "PPTX_MAGIC", "PptBeautifyClient"]
