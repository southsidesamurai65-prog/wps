"""pytest 公共夹具：代码现场生成「内容已知」的小样例文件。

学生不实现任何 TODO，这些夹具也能正常生成——「文件里有什么」就是这里的代码。
PDF 用 fitz insert_text 写已知文本；DOCX/PPTX/PNG 用对应库构造。
"""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


@pytest.fixture(scope="session")
def qapp():
    """共享一个 QApplication（测试 TaskRunner 跨线程信号需要事件循环）。"""
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def tmp_output_dir(tmp_path):
    d = tmp_path / "out"
    d.mkdir()
    return d


# ---------- PDF ----------


def _make_pdf(path: Path, pages: list[str]) -> Path:
    doc = fitz.open()
    for txt in pages:
        page = doc.new_page()
        page.insert_text((72, 72), txt)  # 文本即 spec 断言对象
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def sample_pdf(tmp_path):
    return _make_pdf(tmp_path / "one.pdf", ["Hello PDF World"])


@pytest.fixture
def sample_pdf_multipage(tmp_path):
    return _make_pdf(tmp_path / "multi.pdf", ["PAGE_A", "PAGE_B", "PAGE_C"])


# ---------- DOCX ----------


@pytest.fixture
def sample_docx(tmp_path):
    """token {{COMPANY}} 在单 run 内，便于 replace 简单档。"""
    p = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("{{COMPANY}} is great", style="Title")
    doc.add_paragraph("Second paragraph.")
    doc.save(str(p))
    return p


@pytest.fixture
def sample_docx_multirun(tmp_path):
    """一段 3 个 run，token {{WHO}} 在中间 run——验证替换后 run 数不变。"""
    p = tmp_path / "multirun.docx"
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Hello ")
    para.add_run("{{WHO}}")
    para.add_run(" !")
    doc.save(str(p))
    return p


# ---------- 图片 ----------


@pytest.fixture
def sample_png(tmp_path):
    p = tmp_path / "red.png"
    Image.new("RGB", (100, 100), (255, 0, 0)).save(str(p))
    return p


@pytest.fixture
def sample_pngs(tmp_path):
    colors = [(255, 0, 0), (0, 255, 0), (0, 0, 255)]
    paths = []
    for i, color in enumerate(colors):
        p = tmp_path / f"c{i}.png"
        Image.new("RGB", (50, 50), color).save(str(p))
        paths.append(p)
    return paths


# ---------- PPTX ----------


@pytest.fixture
def sample_pptx(tmp_path, sample_png):
    """slide1 = Title Slide（标题 T1），slide2 = Title and Content（标题 T2 + 1 PNG）。"""
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
    s1.shapes.title.text = "T1"
    s1.placeholders[1].text = "sub"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    s2.shapes.title.text = "T2"
    tf = s2.placeholders[1].text_frame
    tf.text = "b1"
    p2 = tf.add_paragraph()
    p2.text = "b2"
    s2.shapes.add_picture(str(sample_png), Inches(1), Inches(2))
    prs.save(str(p))
    return p
