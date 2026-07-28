"""PPT 本地渲染器 spec（shape 级，纯函数无网络）。

给定固定 shape 级 spec → ``render_beautified_deck`` → 产物可打开、shape 齐全
（rect/rounded_rect/oval/textbox/image）、accent 色填上、图片插入、画布尺寸对、bg override 生效。
"""

from __future__ import annotations

import io

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from wps_tool.services.ppt_renderer import render_beautified_deck


def _png_bytes(color=(0, 200, 0)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), color).save(buf, "PNG")
    return buf.getvalue()


SPEC = {
    "slide_size": {"width": 13.333, "height": 7.5},
    "theme": {"accent": "#1F4E79", "bg": "#FFFFFF"},
    "slides": [
        {
            "bg": "#F0F4FA",
            "shapes": [
                {"type": "rect", "left": 0, "top": 0, "width": 13.333, "height": 1.4,
                 "fill": "#1F4E79", "line": None, "text": "季度汇报", "font_size": 40,
                 "font_color": "#FFFFFF", "bold": True, "align": "left",
                 "valign": "middle", "margin_left": 0.5},
                {"type": "rounded_rect", "left": 0.6, "top": 2, "width": 3, "height": 1.2,
                 "fill": "#E8EEF7", "line": "#1F4E79", "text": "标签"},
                {"type": "oval", "left": 4, "top": 2, "width": 1.5, "height": 1.5,
                 "fill": "#1F4E79"},
                {"type": "image", "image_id": "img_1", "left": 7.6, "top": 2,
                 "width": 5, "height": 4.5},
                {"type": "textbox", "left": 0.6, "top": 4, "width": 6.5, "height": 2.5,
                 "valign": "top",
                 "paragraphs": [
                     {"text": "收入 +18%", "bullet": True, "font_size": 22, "color": "#222222"},
                     {"text": "成本 -4%", "bullet": True},
                 ]},
                {"type": "weird", "left": 0.6, "top": 6.6, "width": 6, "height": 0.6,
                 "text": "兜底"},
            ],
        }
    ],
}


def _shapes(prs):
    return [sh for sl in prs.slides for sh in sl.shapes]


def _has_solid_fill(obj, hex_color: str) -> bool:
    """``obj``（shape 或 slide.background）的 solid fill 是否 == hex_color。"""
    fill = obj.fill
    if fill.type != MSO_FILL.SOLID:
        return False
    try:
        return fill.fore_color.rgb == RGBColor.from_string(hex_color.lstrip("#"))
    except (AttributeError, TypeError):
        return False


def test_renders_shape_spec(tmp_output_dir):
    out = tmp_output_dir / "b.pptx"
    render_beautified_deck(SPEC, str(out), images={"img_1": _png_bytes()})

    assert out.read_bytes()[:4] == b"PK\x03\x04"
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    # 画布尺寸
    assert abs(prs.slide_width / 914400 - 13.333) < 0.01
    assert abs(prs.slide_height / 914400 - 7.5) < 0.01

    shapes = _shapes(prs)
    # 图片插入
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in shapes), "图片没插进去"
    # rect / oval 的 accent 色填上了
    assert any(_has_solid_fill(sh, "1F4E79") for sh in shapes), "accent 色没填上"
    # rounded_rect 的描边色 == accent
    labeled = [
        sh for sh in shapes
        if sh.has_text_frame and sh.text_frame.text == "标签"
    ]
    assert labeled, "标签 shape 没渲染"
    assert labeled[0].line.color.rgb == RGBColor.from_string("1F4E79")
    # autoshape 子型齐全（rect / rounded_rect / oval）
    autos = {str(sh.auto_shape_type) for sh in shapes
             if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE}
    assert any("RECTANGLE" in a for a in autos), autos
    assert any("ROUNDED_RECTANGLE" in a for a in autos), autos
    assert any("OVAL" in a for a in autos), autos
    # 文本在场（含未知 type 兜底成 textbox）
    blob = "\n".join(sh.text_frame.text for sh in shapes if sh.has_text_frame)
    for needle in ("季度汇报", "收入 +18%", "成本 -4%", "标签", "兜底"):
        assert needle in blob, f"缺 {needle!r}"
    # bg override 生效
    assert _has_solid_fill(prs.slides[0].background, "F0F4FA")


def test_image_missing_id_skipped(tmp_output_dir):
    """image_id 在 images 字典里缺失 → 跳过该图，不崩、其它 shape 照渲染。"""
    spec = {
        "slides": [
            {"shapes": [
                {"type": "image", "image_id": "nope", "left": 1, "top": 1,
                 "width": 2, "height": 2},
                {"type": "textbox", "left": 0, "top": 0, "width": 5, "height": 1,
                 "text": "ok"},
            ]}
        ]
    }
    out = tmp_output_dir / "m.pptx"
    render_beautified_deck(spec, str(out))  # 不传 images

    prs = Presentation(str(out))
    shapes = _shapes(prs)
    assert not any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in shapes)
    assert any(sh.has_text_frame and "ok" in sh.text_frame.text for sh in shapes)


def test_theme_defaults_when_missing(tmp_output_dir):
    """spec 无 theme → 用默认 accent/bg，照样产出合法 deck、accent 填上。"""
    spec = {"slides": [{"shapes": [
        {"type": "rect", "left": 0, "top": 0, "width": 5, "height": 1,
         "text": "X"},
    ]}]}
    out = tmp_output_dir / "d.pptx"
    render_beautified_deck(spec, str(out))
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    assert any(_has_solid_fill(sh, "1F4E79") for sh in _shapes(prs))  # 默认 accent
