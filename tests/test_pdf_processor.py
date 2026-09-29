"""PDF 处理器 spec（Layer a）。

未实现时这些测试因 NotImplementedError 红——这就是 spec 看板。
实现 merge/split/rotate/extract_pages/to_images/extract_text 后逐一转绿。
"""

from __future__ import annotations

from pathlib import Path

import fitz
from PIL import Image
from pypdf import PdfReader

from wps_tool.processors.pdf_processor import (
    add_page_numbers,
    add_watermark,
    decrypt_pdf,
    encrypt_pdf,
    extract_pdf_pages,
    extract_pdf_text,
    merge_pdfs,
    pdf_to_docx,
    pdf_to_images,
    rotate_pdf,
    split_pdf,
)


def test_extract_pdf_text(sample_pdf):
    assert "Hello PDF World" in extract_pdf_text(str(sample_pdf))


def test_merge_pdfs(sample_pdf, sample_pdf_multipage, tmp_output_dir):
    out = tmp_output_dir / "merged.pdf"
    merge_pdfs([str(sample_pdf_multipage), str(sample_pdf)], str(out))
    assert out.exists()
    assert len(PdfReader(str(out)).pages) == 3 + 1  # 3 + 1


def test_split_pdf(sample_pdf_multipage, tmp_output_dir):
    paths = split_pdf(str(sample_pdf_multipage), str(tmp_output_dir))
    assert len(paths) == 3
    assert all(Path(p).name == f"page_{i}.pdf" for i, p in enumerate(paths, start=1))
    # 每个拆分文件应恰好 1 页——抓「共用 writer 导致页累积」的错
    # （错的实现会让 page_2 含 2 页、page_3 含 3 页，但文件名/数量仍对）。
    for p in paths:
        assert len(PdfReader(p).pages) == 1, f"{p} 应只有 1 页"


def test_rotate_pdf(sample_pdf_multipage, tmp_output_dir):
    out = tmp_output_dir / "rotated.pdf"
    rotate_pdf(str(sample_pdf_multipage), str(out), angle=90)
    assert out.exists()
    assert len(PdfReader(str(out)).pages) == 3


def test_extract_pdf_pages(sample_pdf_multipage, tmp_output_dir):
    out = tmp_output_dir / "pages.pdf"
    extract_pdf_pages(str(sample_pdf_multipage), [0, 2], str(out))
    assert out.exists()
    assert len(PdfReader(str(out)).pages) == 2


def test_pdf_to_images(sample_pdf_multipage, tmp_output_dir):
    paths = pdf_to_images(str(sample_pdf_multipage), str(tmp_output_dir), zoom=1.0)
    assert len(paths) == 3
    for p in paths:
        assert Image.open(p).size[0] > 0


def test_extract_pdf_text_multipage(sample_pdf_multipage):
    text = extract_pdf_text(str(sample_pdf_multipage))
    for page_text in ("PAGE_A", "PAGE_B", "PAGE_C"):
        assert page_text in text


def test_encrypt_and_decrypt_pdf(sample_pdf, tmp_output_dir):
    enc = tmp_output_dir / "enc.pdf"
    encrypt_pdf(str(sample_pdf), str(enc), user_password="pw")
    assert PdfReader(str(enc)).is_encrypted

    dec = tmp_output_dir / "dec.pdf"
    decrypt_pdf(str(enc), str(dec), "pw")
    reader = PdfReader(str(dec))
    assert not reader.is_encrypted
    assert "Hello PDF World" in reader.pages[0].extract_text()


def test_add_watermark(sample_pdf_multipage, tmp_output_dir):
    out = tmp_output_dir / "wm.pdf"
    add_watermark(str(sample_pdf_multipage), str(out), "CONFIDENTIAL")
    assert out.exists()
    doc = fitz.open(str(out))
    assert len(doc) == 3
    assert "CONFIDENTIAL" in doc[0].get_text()
    doc.close()


def test_add_page_numbers(sample_pdf_multipage, tmp_output_dir):
    out = tmp_output_dir / "numbered.pdf"
    add_page_numbers(str(sample_pdf_multipage), str(out))
    doc = fitz.open(str(out))
    assert len(doc) == 3
    text = doc[2].get_text()
    assert "3" in text
    doc.close()


def test_pdf_to_docx(sample_pdf_multipage, tmp_output_dir):
    from docx import Document

    out = tmp_output_dir / "converted.docx"
    pdf_to_docx(str(sample_pdf_multipage), str(out))
    assert out.exists()
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    for page_text in ("PAGE_A", "PAGE_B", "PAGE_C"):
        assert page_text in text


def test_pdf_to_docx_skips_control_chars(tmp_path, tmp_output_dir):
    """回归：PDF 文本层含 \\x00/\\x01 等控制字符时不能抛 XML 兼容错误。"""
    from docx import Document

    src = tmp_path / "ctrl.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "A\x01B")
    page.insert_text((72, 100), "C\x00D")
    doc.save(str(src))
    doc.close()

    out = tmp_output_dir / "ctrl.docx"
    pdf_to_docx(str(src), str(out))
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "AB" in text and "CD" in text

