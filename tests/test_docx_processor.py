"""DOCX 处理器 spec（Layer a）。

replace 分两档：简单档（token 在单 run，断言文本命中）+ 保样式档（多 run，
断言替换后 run 数不变 → 抓「粗暴 paragraph.text= 丢 run」的常见错）。
"""

from __future__ import annotations

from docx import Document

from wps_tool.processors.docx_processor import extract_docx_text, replace_docx_text


def test_extract_docx_text(sample_docx):
    text = extract_docx_text(str(sample_docx))
    assert "{{COMPANY}} is great" in text
    assert "Second paragraph." in text


def test_replace_docx_text_simple(sample_docx, tmp_output_dir):
    out = tmp_output_dir / "replaced.docx"
    replace_docx_text(str(sample_docx), str(out), {"{{COMPANY}}": "ACME"})
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "ACME is great" in texts
    assert not any("{{COMPANY}}" in t for t in texts)


def test_replace_docx_text_preserves_runs(sample_docx_multirun, tmp_output_dir):
    """替换后 run 数应不变——用 paragraph.text= 会让 run 数变 1，此处会失败。"""
    out = tmp_output_dir / "replaced.docx"
    replace_docx_text(str(sample_docx_multirun), str(out), {"{{WHO}}": "World"})

    doc = Document(str(out))
    # 找到含 "World" 的段
    target = next(p for p in doc.paragraphs if "World" in p.text)
    assert target.text == "Hello World !"
    assert len(target.runs) == 3  # 原始也是 3 个 run：保样式档
