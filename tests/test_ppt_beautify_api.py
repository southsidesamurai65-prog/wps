"""PPT 美化客户端 spec（LLM + 本地渲染，shape 级，离线）。

用 httpx.MockTransport 注入：handler 回一段 canned OpenAI Chat Completions 响应
（shape 级 spec，引用 img_1）。断言：
  - 产物是合法 pptx、2 页、spec 文本在场、**图片被搬到新 deck**（picture shape 在场）；
  - 隐私：请求体**不含** pptx 二进制（PK 头）也**不含**图片字节（PNG 头），
    但含图片占位 image_id 和源文本（只发文本+占位）；
  - 请求打到 /v1/chat/completions、带 Authorization、body 含 model 与 response_format。
"""

from __future__ import annotations

import json

import httpx
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from wps_tool.core.errors import ApiUnavailableError
from wps_tool.services.ppt_beautify_api import (
    PNG_MAGIC,
    PPTX_MAGIC,
    PptBeautifyClient,
)

#: canned shape 级 spec：2 页，对齐 sample_pptx（2 页、1 张图 img_1）。
SPEC = {
    "slide_size": {"width": 10, "height": 7.5},
    "theme": {"accent": "#1F4E79", "bg": "#FFFFFF"},
    "slides": [
        {"bg": "#FFFFFF", "shapes": [
            {"type": "rect", "left": 0, "top": 0, "width": 10, "height": 1.4,
             "fill": "#1F4E79", "text": "重设T1", "font_color": "#FFFFFF",
             "bold": True, "valign": "middle", "margin_left": 0.5},
            {"type": "textbox", "left": 0.6, "top": 2, "width": 8.8, "height": 1,
             "text": "新副标题"},
        ]},
        {"bg": "#FFFFFF", "shapes": [
            {"type": "rect", "left": 0, "top": 0, "width": 10, "height": 1.4,
             "fill": "#1F4E79", "text": "重设T2", "font_color": "#FFFFFF",
             "bold": True, "valign": "middle", "margin_left": 0.5},
            {"type": "image", "image_id": "img_1", "left": 5, "top": 2,
             "width": 4, "height": 4},
            {"type": "textbox", "left": 0.6, "top": 2, "width": 4, "height": 4,
             "valign": "top",
             "paragraphs": [
                 {"text": "新b1", "bullet": True},
                 {"text": "新b2", "bullet": True},
             ]},
        ]},
    ],
}


def _chat_completion_handler(captured: dict):
    def handler(request: httpx.Request) -> httpx.Response:
        captured["request"] = request
        captured["content"] = request.content
        captured["body"] = json.loads(request.content)
        content = json.dumps(SPEC, ensure_ascii=False)
        return httpx.Response(
            200, json={"choices": [{"message": {"content": content}}]}
        )

    return handler


def test_beautify_carries_image_and_sends_text_only(sample_pptx, tmp_output_dir):
    captured: dict = {}
    out = tmp_output_dir / "beautified.pptx"
    transport = httpx.MockTransport(_chat_completion_handler(captured))
    client = PptBeautifyClient("key", model="gpt-4o", transport=transport)

    result = client.beautify_file(str(sample_pptx), str(out), style="business")

    assert result == str(out)
    # 产物合法 pptx、2 页
    assert out.read_bytes().startswith(PPTX_MAGIC)
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    # 图片被搬到新 deck（关键回归点：之前美化后图全丢）
    shapes = [sh for sl in prs.slides for sh in sl.shapes]
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in shapes), "图片没搬过来"
    text = "\n".join(sh.text_frame.text for sh in shapes if sh.has_text_frame)
    assert "重设T1" in text and "重设T2" in text

    # 隐私：请求体不含 pptx 二进制、不含图片字节
    assert PPTX_MAGIC not in captured["content"], "请求体含 pptx 二进制"
    assert PNG_MAGIC not in captured["content"], "请求体含图片字节"
    # 但含图片占位 + 源文本（只发文本+占位）
    assert b"img_1" in captured["content"], "没把图片占位发给 LLM"
    assert b"T1" in captured["content"] and b"T2" in captured["content"]

    # 请求打到 OpenAI Chat Completions，带鉴权，body 含 model 与 response_format
    assert captured["request"].url.path == "/v1/chat/completions"
    assert captured["request"].headers["Authorization"] == "Bearer key"
    assert captured["body"]["model"] == "gpt-4o"
    assert captured["body"]["response_format"] == {"type": "json_object"}
    client.close()


def test_beautify_file_empty_key_raises(tmp_output_dir, sample_pptx):
    client = PptBeautifyClient(
        "", model="gpt-4o",
        transport=httpx.MockTransport(lambda r: httpx.Response(200)),
    )
    with pytest.raises(ApiUnavailableError):
        client.beautify_file(str(sample_pptx), str(tmp_output_dir / "x.pptx"))
    client.close()


def test_unsupported_provider_raises():
    with pytest.raises(ApiUnavailableError):
        PptBeautifyClient("key", model="gpt-4o", provider="gemini")


def test_build_beautify_client_uses_llm_base_url():
    """app.py 装配：LLM_BASE_URL 真传进 client.base_url（不写死）。

    防回归：以后 build_beautify_client 漏传 base_url 时此测试会红——
    会退回 client 构造器的默认 https://api.openai.com，与配置不符。
    """
    from wps_tool.app import build_beautify_client
    from wps_tool.models.settings import Settings

    s = Settings(
        _env_file=None,
        enable_api_upload=True,
        llm_provider="openai",
        llm_api_key="k",
        llm_base_url="https://api.my-relay.example.com",
    )
    client = build_beautify_client(s)
    assert client is not None
    assert (
        str(client.client.base_url).rstrip("/")
        == "https://api.my-relay.example.com"
    )
    client.close()
