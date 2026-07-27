"""PPTX 处理器 spec（Layer a）。"""

from __future__ import annotations

from pptx import Presentation

from wps_tool.processors.ppt_processor import (
    analyze_pptx_structure,
    extract_pptx_images,
    extract_pptx_text,
    replace_pptx_tokens,
)


def test_extract_pptx_text(sample_pptx):
    result = extract_pptx_text(str(sample_pptx))
    by_slide = {r["slide"]: r["texts"] for r in result}
    assert "T1" in by_slide[1]
    assert "T2" in by_slide[2]
    assert "b1" in by_slide[2] and "b2" in by_slide[2]


def test_extract_pptx_images(sample_pptx, tmp_output_dir):
    paths = extract_pptx_images(str(sample_pptx), str(tmp_output_dir))
    assert len(paths) == 1
    assert all(p.endswith((".png", ".jpeg", ".jpg")) for p in paths)


def test_analyze_pptx_structure(sample_pptx):
    result = analyze_pptx_structure(str(sample_pptx))
    by_slide = {r["slide"]: r for r in result}
    assert by_slide[1]["page_type"] == "title"
    assert by_slide[2]["page_type"] == "content"


def test_replace_pptx_tokens(sample_pptx, tmp_output_dir):
    out = tmp_output_dir / "replaced.pptx"
    replace_pptx_tokens(str(sample_pptx), str(out), {"T1": "X1", "T2": "X2"})

    prs = Presentation(str(out))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    assert any("X1" in t for t in texts)
    assert any("X2" in t for t in texts)
    assert not any("T1" in t for t in texts)
